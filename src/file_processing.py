import os

import docx

from Caesar import *
from Vigenere import *


def processing(path, key, cipher, todo, error, key_error, key_entry, path_style, key_style):
    how_coded = ("encrypted", "decrypted")
    if path.get().endswith('\"'):
        path.set(path.get()[1:len(path.get()) - 1])
    file = os.path.abspath(path.get())
    file, ext = os.path.splitext(file)
    if not os.path.exists(path.get()):
        path.set(error)
        path_style.map('MyPath.TEntry', foreground=[('!focus', 'red'), ('focus', 'black')])
    else:
        if '.doc' in ext:
            doc = docx.Document(path.get())
            new_doc = docx.Document()
            if key == '-':
                text = ''
                for paragraph in doc.paragraphs:
                    for sym in paragraph.text:
                        if sym.isalpha:
                            text += sym
                key = sym_frequency(text)
                if key == 'error':
                    key_style.map('MyKey.TEntry', foreground=[('!focus', 'red'), ('focus', 'black')])
                    key_entry.set(key_error)
                    return
            for paragraph in doc.paragraphs:
                new_paragraph = new_doc.add_paragraph()
                for sym in paragraph.text:
                    if sym.isalpha():
                        sym = change(cipher, todo, key, sym)
                    new_paragraph.add_run(sym)
            new_doc.save(file + "_" + how_coded[todo] + ext)
        else:
            with open(path.get(), 'r', encoding='utf-8', newline='', errors='ignore') as old_file:
                if key == '-':
                    text = ''
                    line = old_file.readline()
                    while line != '':
                        for sym in line:
                            if sym.isalpha:
                                text += sym
                        line = old_file.readline()
                    key = sym_frequency(text)
                    if key == 'error':
                        key_style.map('MyKey.TEntry', foreground=[('!focus', 'red'), ('focus', 'black')])
                        key_entry.set(key_error)
                        return
            with open(path.get(), 'r', encoding='utf-8', newline='', errors='ignore') as old_file:
                with open(file + "_" + how_coded[todo] + ext, 'x') as new_file:
                    line = old_file.readline()
                    while line != '':
                        for sym in line:
                            if sym.isalpha():
                                sym = change(cipher, todo, key, sym)
                            new_file.write(sym)
                        new_file.write('\n')
                        line = old_file.readline()
        return file + "_" + how_coded[todo] + ext, 'x'


def change(cipher, todo, key, sym):
    if cipher == 0:
        return caesar(todo, key, sym)
    elif cipher == 1:
        return vigenere(todo, key, sym)
    else:
        return sym
